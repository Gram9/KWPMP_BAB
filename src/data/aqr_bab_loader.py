"""
Loader for AQR's published "Betting Against Beta" factor workbook
(config/aqr_bab.yaml; see docs/02_validation_gates.md Gate 4 for the full
derivation and the anchor values this loader's tests assert against).

Sheet "BAB Factors" has a two-row header block (row 18 groups columns
under labels like "EQUITIES", row 19 has the actual per-country column
names) with data starting row 20. Column position, not header text
matching, is what config/aqr_bab.yaml pins (USA=Y, CAN=E) -- AQR's header
text is a short country code repeated across ~25 columns, not unique
enough to locate reliably by search the way ken_french_loader.py's block
markers are. The loader instead asserts the header ROW says what config
expects at the DATE column, which is the trap this file's own history
already hit once (docs/02_validation_gates.md Gate 4): a column-offset or
header-row shift is the realistic failure mode, and a silent misparse
there would just read some other country's column as USA/CAN.

Values on this sheet are already decimal returns; do not divide by 100.
"""

import datetime as dt
from pathlib import Path

import openpyxl
import polars as pl
import yaml
from openpyxl.utils import column_index_from_string

PROJECT_ROOT = Path(__file__).resolve().parents[2]
CONFIG_PATH = PROJECT_ROOT / "config" / "aqr_bab.yaml"

with CONFIG_PATH.open() as f:
    _CONFIG = yaml.safe_load(f)

WORKBOOK_PATH = PROJECT_ROOT / _CONFIG["path"]
CACHE_PATH = PROJECT_ROOT / _CONFIG["cache_path"]
_SHEET_NAME = _CONFIG["sheet"]
_HEADER_ROW = _CONFIG["header_row"]
_DATA_FIRST_ROW = _HEADER_ROW + 1
_DATE_COL = column_index_from_string(_CONFIG["date_column"])
_USA_COL = column_index_from_string(_CONFIG["columns"]["usa"])
_CAN_COL = column_index_from_string(_CONFIG["columns"]["can"])
assert _CONFIG["units"] == "decimal", (
    f"{CONFIG_PATH}: expected units: decimal -- loader does not implement "
    "any other unit conversion for this sheet."
)


def load_aqr_bab_factors() -> pl.DataFrame:
    """AQR's published BAB factor, US and Canada legs: date, usa_ret,
    can_ret (decimal monthly returns, both nullable -- CAN starts in
    1987, well after USA's 1930 start). Cached to CACHE_PATH; re-parses
    the source workbook only when it's missing or older than the source
    workbook's mtime, so a re-download is picked up automatically rather
    than silently served a stale cache."""
    if CACHE_PATH.exists() and CACHE_PATH.stat().st_mtime >= WORKBOOK_PATH.stat().st_mtime:
        return pl.read_parquet(CACHE_PATH)

    result = _parse_aqr_bab_workbook()
    CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
    result.write_parquet(CACHE_PATH)
    return result


def _parse_aqr_bab_workbook() -> pl.DataFrame:
    wb = openpyxl.load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    ws = wb[_SHEET_NAME]

    header_date_cell = ws.cell(row=_HEADER_ROW, column=_DATE_COL).value
    assert header_date_cell == "DATE", (
        f"{WORKBOOK_PATH}: expected 'DATE' at sheet '{_SHEET_NAME}' row "
        f"{_HEADER_ROW}, column {_DATE_COL}, got {header_date_cell!r} -- "
        "header row may have shifted since config/aqr_bab.yaml was "
        "written against it; refusing to silently misparse."
    )

    dates: list = []
    usa_rets: list[float | None] = []
    can_rets: list[float | None] = []

    for row in ws.iter_rows(
        min_row=_DATA_FIRST_ROW,
        min_col=1,
        max_col=max(_DATE_COL, _USA_COL, _CAN_COL),
        values_only=True,
    ):
        date_val = row[_DATE_COL - 1]
        if date_val is None:
            continue
        if isinstance(date_val, str):
            # Suppression reason: source dates are plain MM/DD/YYYY
            # calendar dates with no timezone semantics -- naive
            # datetime is intentional and correct here.
            date_val = dt.datetime.strptime(date_val, "%m/%d/%Y").date()  # noqa: DTZ007
        elif isinstance(date_val, dt.datetime):
            date_val = date_val.date()
        dates.append(date_val)
        usa_val = row[_USA_COL - 1]
        can_val = row[_CAN_COL - 1]
        assert usa_val is None or isinstance(usa_val, (int, float)), (
            f"{WORKBOOK_PATH}: non-numeric USA cell {usa_val!r} at row "
            f"with date {date_val!r} -- config/aqr_bab.yaml declares "
            "units: decimal."
        )
        assert can_val is None or isinstance(can_val, (int, float)), (
            f"{WORKBOOK_PATH}: non-numeric CAN cell {can_val!r} at row "
            f"with date {date_val!r} -- config/aqr_bab.yaml declares "
            "units: decimal."
        )
        usa_rets.append(float(usa_val) if usa_val is not None else None)
        can_rets.append(float(can_val) if can_val is not None else None)

    return pl.DataFrame(
        {"date": dates, "usa_ret": usa_rets, "can_ret": can_rets},
        schema={"date": pl.Date, "usa_ret": pl.Float64, "can_ret": pl.Float64},
    )
