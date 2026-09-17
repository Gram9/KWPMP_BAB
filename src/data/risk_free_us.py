"""
US risk-free rate loaders, per config/risk_free.yaml's `us` block (full
reasoning lives in that file -- this module honours those decisions
rather than re-deriving them).

Two independent paths, NOT interchangeable:

- load_us_rf_ken_french(): Ken French's monthly RF from the standard
  3-factor CSV. Percent units (divide by 100). The file holds TWO
  tables -- monthly data, then a blank line, then an "Annual Factors"
  block keyed by bare 4-digit year. A naive read_csv parses those years
  as dates and silently appends annual returns onto the monthly series;
  this loader stops at the configured last monthly line and asserts the
  blank-line/annual-marker boundary is exactly where config says, so a
  future re-download with a shifted table boundary fails loudly instead
  of silently absorbing annual rows.
- load_us_rf_aqr_gate4(): AQR's own RF sheet from the BAB workbook.
  Gate 4 ONLY -- AQR built their published BAB series with this RF, so
  it is the faithful comparison input for that gate. Close to but not
  identical with Ken French's; the two must never be used
  interchangeably (see config/risk_free.yaml).
"""

import datetime as dt
from pathlib import Path

import openpyxl
import polars as pl
import yaml
from openpyxl.utils import column_index_from_string

PROJECT_ROOT = Path(__file__).resolve().parents[2]
CONFIG_PATH = PROJECT_ROOT / "config" / "risk_free.yaml"

with CONFIG_PATH.open() as f:
    _CONFIG = yaml.safe_load(f)

_US_CONFIG = _CONFIG["us"]
KEN_FRENCH_PATH = PROJECT_ROOT / _US_CONFIG["path"]
_KEN_FRENCH_COLUMN = _US_CONFIG["column"]
_KEN_FRENCH_HEADER_LINE = _US_CONFIG["monthly_header_line"]
_KEN_FRENCH_LAST_LINE = _US_CONFIG["monthly_last_line"]
assert _US_CONFIG["units"] == "percent", (
    f"{CONFIG_PATH}: us.units expected 'percent' -- loader divides by "
    "100 unconditionally and does not implement any other conversion."
)

_ANNUAL_BLOCK_MARKER = "Annual Factors"

_GATE4_CONFIG = _US_CONFIG["gate4_override"]
GATE4_WORKBOOK_PATH = PROJECT_ROOT / _GATE4_CONFIG["path"]
_GATE4_SHEET_NAME = _GATE4_CONFIG["sheet"]
_GATE4_HEADER_ROW = _GATE4_CONFIG["header_row"]
_GATE4_DATE_COL = column_index_from_string(_GATE4_CONFIG["date_column"])
_GATE4_VALUE_COL = column_index_from_string(_GATE4_CONFIG["value_column"])
assert _GATE4_CONFIG["units"] == "decimal", (
    f"{CONFIG_PATH}: us.gate4_override.units expected 'decimal' -- loader "
    "does not implement any other conversion for this sheet."
)


def load_us_rf_ken_french() -> pl.DataFrame:
    """Ken French monthly RF: date (month-end), rf (decimal). Parses only
    the monthly table -- the file's separate annual-factors block is
    never read."""
    lines = KEN_FRENCH_PATH.read_text().splitlines()

    header_idx = _KEN_FRENCH_HEADER_LINE
    blank_idx = _KEN_FRENCH_LAST_LINE + 1
    marker_idx = blank_idx + 1
    assert len(lines) > marker_idx, (
        f"{KEN_FRENCH_PATH}: expected at least {marker_idx + 1} lines "
        f"(header at {header_idx}, blank separator at {blank_idx}, "
        f"'{_ANNUAL_BLOCK_MARKER}' at {marker_idx}), got {len(lines)} -- "
        "file is shorter than config/risk_free.yaml's monthly_last_line "
        "expects; refusing to silently misparse."
    )

    columns = [c.strip() for c in lines[header_idx].split(",")]
    rf_col_idx = columns.index(_KEN_FRENCH_COLUMN)

    assert not lines[blank_idx].strip(), (
        f"{KEN_FRENCH_PATH}: expected line {blank_idx} (0-indexed) to be "
        "blank, separating the monthly table from the annual-factors "
        "block -- file layout may have shifted since "
        "config/risk_free.yaml's monthly_last_line was set; refusing to "
        "silently misparse and possibly absorb annual rows."
    )
    assert _ANNUAL_BLOCK_MARKER in lines[marker_idx], (
        f"{KEN_FRENCH_PATH}: expected line {marker_idx} (0-indexed) to "
        f"contain '{_ANNUAL_BLOCK_MARKER}' -- file layout may have "
        "shifted since config/risk_free.yaml's monthly_last_line was "
        "set; refusing to silently misparse."
    )

    dates: list[dt.date] = []
    rfs: list[float] = []
    for line in lines[header_idx + 1 : _KEN_FRENCH_LAST_LINE + 1]:
        fields = [f.strip() for f in line.split(",")]
        yyyymm = fields[0]
        # Suppression reason: source keys are plain YYYYMM calendar
        # months with no timezone semantics -- naive datetime is correct.
        first_of_month = dt.datetime.strptime(yyyymm, "%Y%m")  # noqa: DTZ007
        next_month = (first_of_month.replace(day=1) + dt.timedelta(days=32)).replace(day=1)
        month_end = (next_month - dt.timedelta(days=1)).date()
        dates.append(month_end)
        rfs.append(float(fields[rf_col_idx]) / 100.0)

    return pl.DataFrame(
        {"date": dates, "rf": rfs},
        schema={"date": pl.Date, "rf": pl.Float64},
    )


def load_us_rf_aqr_gate4() -> pl.DataFrame:
    """AQR's own RF sheet: date (month-end), rf (decimal). Gate 4 use
    only -- see module docstring."""
    wb = openpyxl.load_workbook(GATE4_WORKBOOK_PATH, read_only=True, data_only=True)
    ws = wb[_GATE4_SHEET_NAME]

    header_date_cell = ws.cell(row=_GATE4_HEADER_ROW, column=_GATE4_DATE_COL).value
    assert header_date_cell == "DATE", (
        f"{GATE4_WORKBOOK_PATH}: expected 'DATE' at sheet "
        f"'{_GATE4_SHEET_NAME}' row {_GATE4_HEADER_ROW}, column "
        f"{_GATE4_DATE_COL}, got {header_date_cell!r} -- header row may "
        "have shifted since config/risk_free.yaml was written against "
        "it; refusing to silently misparse."
    )

    dates: list[dt.date] = []
    rfs: list[float | None] = []
    for row in ws.iter_rows(
        min_row=_GATE4_HEADER_ROW + 1,
        min_col=1,
        max_col=max(_GATE4_DATE_COL, _GATE4_VALUE_COL),
        values_only=True,
    ):
        date_val = row[_GATE4_DATE_COL - 1]
        if date_val is None:
            continue
        if isinstance(date_val, str):
            # Suppression reason: source dates are plain MM/DD/YYYY
            # calendar dates with no timezone semantics -- naive
            # datetime is intentional and correct here.
            date_val = dt.datetime.strptime(date_val, "%m/%d/%Y").date()  # noqa: DTZ007
        elif isinstance(date_val, dt.datetime):
            date_val = date_val.date()
        assert isinstance(date_val, dt.date), (
            f"{GATE4_WORKBOOK_PATH}: unexpected date cell type "
            f"{type(date_val)!r} ({date_val!r})."
        )
        rf_val = row[_GATE4_VALUE_COL - 1]
        assert rf_val is None or isinstance(rf_val, (int, float)), (
            f"{GATE4_WORKBOOK_PATH}: non-numeric RF cell {rf_val!r} at "
            f"row with date {date_val!r} -- "
            "config/risk_free.yaml declares gate4_override units: decimal."
        )
        dates.append(date_val)
        rfs.append(float(rf_val) if rf_val is not None else None)

    return pl.DataFrame(
        {"date": dates, "rf": rfs},
        schema={"date": pl.Date, "rf": pl.Float64},
    )
