"""Tests for src/data/aqr_bab_loader.py -- AQR's published BAB factor
workbook (config/aqr_bab.yaml, docs/02_validation_gates.md Gate 4).

Two kinds of test here, deliberately kept separate:

1. Parsing-logic tests against a small synthetic workbook that replicates
   the real file's exact layout (header row 19, data from row 20, DATE in
   column A, CAN in column E, USA in column Y) -- these catch a
   column-offset or header-row-detection bug in the loader itself.
2. Anchor tests against the REAL workbook in data/raw/ -- these are the
   ones that actually catch a column-offset error in the source file
   (the realistic failure mode per docs/02_validation_gates.md), and are
   the numbers CLAUDE.md requires as assertions, not eyeballing.
"""

import datetime

import openpyxl
import polars as pl
import pytest

from src.data import aqr_bab_loader

_HEADER_ROW = 19
_USA_COL = 25  # Y
_CAN_COL = 5  # E


def _write_fixture_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BAB Factors"
    ws.cell(row=18, column=2, value="EQUITIES")
    ws.cell(row=_HEADER_ROW, column=1, value="DATE")
    ws.cell(row=_HEADER_ROW, column=_CAN_COL, value="CAN")
    ws.cell(row=_HEADER_ROW, column=_USA_COL, value="USA")
    rows = [
        ("12/31/1930", None, -0.000557986),
        ("01/31/1987", None, 0.01),
        ("02/28/1987", 0.02, 0.011),
        ("03/31/1987", 0.03, 0.012),
    ]
    for offset, (date_str, can_val, usa_val) in enumerate(rows):
        r = 20 + offset
        ws.cell(row=r, column=1, value=date_str)
        ws.cell(row=r, column=_CAN_COL, value=can_val)
        ws.cell(row=r, column=_USA_COL, value=usa_val)
    wb.save(path)


@pytest.fixture
def fixture_workbook(tmp_path, monkeypatch):
    path = tmp_path / "aqr_fixture.xlsx"
    _write_fixture_workbook(path)
    monkeypatch.setattr(aqr_bab_loader, "WORKBOOK_PATH", path)
    monkeypatch.setattr(aqr_bab_loader, "CACHE_PATH", tmp_path / "aqr_fixture_cache.parquet")
    return path


def test_load_aqr_bab_factors_reads_usa_and_can_columns(fixture_workbook):
    result = aqr_bab_loader.load_aqr_bab_factors()

    assert set(result.columns) == {"date", "usa_ret", "can_ret"}
    assert result.schema["date"] == pl.Date
    assert result.height == 4


def test_load_aqr_bab_factors_values_are_decimal_not_percent(fixture_workbook):
    result = aqr_bab_loader.load_aqr_bab_factors()

    row = result.filter(pl.col("date") == pl.date(1930, 12, 31))
    assert row.height == 1
    # Fixture value is -0.000557986 -- must NOT be divided by 100 (would
    # be -0.00000557986, which is what a wrong /100 conversion produces).
    assert row["usa_ret"][0] == pytest.approx(-0.000557986, rel=1e-9)


def test_load_aqr_bab_factors_can_has_earlier_nulls_than_usa(fixture_workbook):
    result = aqr_bab_loader.load_aqr_bab_factors()

    early = result.filter(pl.col("date") == pl.date(1930, 12, 31))
    assert early["can_ret"][0] is None
    assert early["usa_ret"][0] is not None

    first_can_row = result.filter(pl.col("can_ret").is_not_null()).sort("date").head(1)
    assert first_can_row["date"][0] == datetime.date(1987, 2, 28)


def test_load_aqr_bab_factors_raises_if_header_row_shifted(tmp_path, monkeypatch):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BAB Factors"
    # Header written one row off from the configured header_row=19 --
    # this is exactly the column-offset/header-shift failure mode the
    # roadmap calls out as the realistic risk.
    ws.cell(row=_HEADER_ROW + 1, column=1, value="DATE")
    ws.cell(row=_HEADER_ROW + 1, column=_USA_COL, value="USA")
    path = tmp_path / "shifted.xlsx"
    wb.save(path)
    monkeypatch.setattr(aqr_bab_loader, "WORKBOOK_PATH", path)

    with pytest.raises(AssertionError, match="DATE"):
        aqr_bab_loader.load_aqr_bab_factors()


def test_load_aqr_bab_factors_writes_parquet_cache(fixture_workbook):
    assert not aqr_bab_loader.CACHE_PATH.exists()

    result = aqr_bab_loader.load_aqr_bab_factors()

    assert aqr_bab_loader.CACHE_PATH.exists()
    cached = pl.read_parquet(aqr_bab_loader.CACHE_PATH)
    assert cached.equals(result)


def test_load_aqr_bab_factors_reuses_cache_when_source_unchanged(fixture_workbook):
    first = aqr_bab_loader.load_aqr_bab_factors()
    cache_mtime_after_first_call = aqr_bab_loader.CACHE_PATH.stat().st_mtime_ns

    second = aqr_bab_loader.load_aqr_bab_factors()

    # Cache file must not have been rewritten -- the workbook wasn't
    # touched, so the second call should read the cache rather than
    # re-parsing and re-writing it.
    assert aqr_bab_loader.CACHE_PATH.stat().st_mtime_ns == cache_mtime_after_first_call
    assert second.equals(first)


def test_load_aqr_bab_factors_reparses_when_source_is_newer_than_cache(
    fixture_workbook, tmp_path
):
    first = aqr_bab_loader.load_aqr_bab_factors()
    assert first.height == 4

    # Rewrite the source workbook with one extra row, then bump its
    # mtime past the cache's -- the loader must detect this and
    # re-parse rather than silently serving the stale cached frame.
    wb = openpyxl.load_workbook(fixture_workbook)
    ws = wb["BAB Factors"]
    ws.cell(row=24, column=1, value="04/30/1987")
    ws.cell(row=24, column=_USA_COL, value=0.013)
    wb.save(fixture_workbook)
    new_mtime = aqr_bab_loader.CACHE_PATH.stat().st_mtime + 5
    import os

    os.utime(fixture_workbook, (new_mtime, new_mtime))

    second = aqr_bab_loader.load_aqr_bab_factors()

    assert second.height == 5


class TestRealWorkbookAnchors:
    """Anchors against the real file -- see docs/02_validation_gates.md
    Gate 4. These are the assertions CLAUDE.md requires in place of
    eyeballing the output; a column-offset or header-row error in the
    real workbook would move these values, not just the fixture's."""

    @staticmethod
    @pytest.fixture(scope="class")
    def result():
        return aqr_bab_loader.load_aqr_bab_factors()

    def test_usa_first_row(self, result):
        row = result.filter(pl.col("date") == pl.date(1930, 12, 31))
        assert row.height == 1
        assert row["usa_ret"][0] == pytest.approx(-0.000557986, rel=1e-6)

    def test_usa_last_row(self, result):
        row = result.filter(pl.col("date") == pl.date(2026, 6, 30))
        assert row.height == 1
        assert row["usa_ret"][0] == pytest.approx(0.0224646878, rel=1e-6)

    def test_usa_non_null_count(self, result):
        assert result["usa_ret"].drop_nulls().len() == 1147

    def test_can_first_non_null(self, result):
        first = result.filter(pl.col("can_ret").is_not_null()).sort("date").head(1)
        assert first["date"][0] == datetime.date(1987, 2, 28)

    def test_can_non_null_count(self, result):
        assert result["can_ret"].drop_nulls().len() == 473
